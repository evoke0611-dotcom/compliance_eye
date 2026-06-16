"""
Fills the ISMS POLICY Template.docx with data collected by Alex
and returns a BytesIO object ready for download.
"""

import io
import datetime
import os
from copy import deepcopy

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "ISMS POLICY Template.docx")


# ── helpers ───────────────────────────────────────────────────────────────

def _v(data: dict, key: str, fallback: str = "—") -> str:
    """Safe value getter."""
    return str(data.get(key) or fallback).strip() or fallback


def _set_cell(cell, text: str):
    """Replace every run in a cell with plain text."""
    for para in cell.paragraphs:
        for run in para.runs:
            run.text = ""
    if cell.paragraphs:
        cell.paragraphs[0].clear()
        cell.paragraphs[0].add_run(text)
    else:
        cell.add_paragraph(text)


def _fill_table(table, mapping: dict):
    """
    Walk every cell in a table. If cell text matches a key in *mapping*,
    replace it with the value.
    """
    for row in table.rows:
        for cell in row.cells:
            stripped = cell.text.strip()
            if stripped in mapping:
                _set_cell(cell, mapping[stripped])


def _replace_placeholder(doc, placeholder: str, value: str):
    """Replace a text placeholder across all paragraphs and table cells."""
    for para in doc.paragraphs:
        if placeholder in para.text:
            for run in para.runs:
                if placeholder in run.text:
                    run.text = run.text.replace(placeholder, value)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if placeholder in para.text:
                        for run in para.runs:
                            if placeholder in run.text:
                                run.text = run.text.replace(placeholder, value)


# ── main builder ──────────────────────────────────────────────────────────

def build_isms_policy_docx(collected: dict) -> io.BytesIO:
    """
    Args:
        collected: flat dict of all answers (from st.session_state.isms_collected)
                   plus lists: collected['risks'], collected['assets'], collected['backups']

    Returns:
        BytesIO of the filled DOCX.
    """
    if not os.path.exists(TEMPLATE_PATH):
        raise FileNotFoundError(f"Template not found: {TEMPLATE_PATH}")

    doc = Document(TEMPLATE_PATH)

    d = collected  # shorthand
    today = datetime.datetime.now().strftime("%Y-%m-%d")

    # ── 1. Simple placeholder replacements ───────────────────────────────
    placeholders = {
        "[Organization Name (Security Division)]": _v(d, "org_name"),
        "[Dummy Office Location 1]":               _v(d, "org_location"),
        "itcindia@gmail.com":                      _v(d, "org_email"),
    }
    for ph, val in placeholders.items():
        _replace_placeholder(doc, ph, val)

    # ── 2. Table 0 — Cover / Header ───────────────────────────────────────
    # Cells: ORGANISATION NAME, AUDITOR NAME, AUDITING DATE, CLASSIFICATION
    if len(doc.tables) > 0:
        t = doc.tables[0]
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            for ci, cell in enumerate(row.cells):
                txt = cell.text.strip()
                if txt == "ORGANISATION NAME":
                    # value is next cell
                    if ci + 1 < len(row.cells):
                        _set_cell(row.cells[ci + 1], _v(d, "org_name"))
                elif txt == "AUDITOR NAME":
                    if ci + 1 < len(row.cells):
                        _set_cell(row.cells[ci + 1], _v(d, "auditor_name"))
                elif txt == "AUDITING DATE":
                    if ci + 1 < len(row.cells):
                        _set_cell(row.cells[ci + 1], _v(d, "audit_date", today))

    # ── 3. Table 2 — Audit Details ────────────────────────────────────────
    if len(doc.tables) > 2:
        t = doc.tables[2]
        label_to_value = {
            "Audit Type (Internal / External)": _v(d, "audit_type"),
            "Audit Date":     _v(d, "audit_date", today),
            "Auditor Name":   _v(d, "auditor_name"),
            "Department Audited": _v(d, "dept_audited"),
        }
        for row in t.rows:
            for ci, cell in enumerate(row.cells):
                if cell.text.strip() in label_to_value and ci + 1 < len(row.cells):
                    _set_cell(row.cells[ci + 1], label_to_value[cell.text.strip()])

    # ── 4. Table 3 — Organisation Details ────────────────────────────────
    if len(doc.tables) > 3:
        t = doc.tables[3]
        label_to_value = {
            "Organization Name":  _v(d, "org_name"),
            "Organization E-mail": _v(d, "org_email"),
            "Nature of Business": _v(d, "org_business"),
            "Office Location(s)": _v(d, "org_location"),
            "Number of Employees": _v(d, "org_employees"),
            "IT Infrastructure (On-Premise/ Cloud)": _v(d, "org_it_infra"),
        }
        for row in t.rows:
            for ci, cell in enumerate(row.cells):
                if cell.text.strip() in label_to_value and ci + 1 < len(row.cells):
                    _set_cell(row.cells[ci + 1], label_to_value[cell.text.strip()])

    # ── 5. Table 4 — Management & ISMS Team ──────────────────────────────
    if len(doc.tables) > 4:
        t = doc.tables[4]
        # Template rows: CEO/Director | ISMS Manager | HR Head | System Administrator | Technical Head
        # Cols: Role | Name | Department | Responsibilities | Contact
        role_map = {
            "CEO/ Director":        ("ceo_name",      "ceo_dept",      "Overall ISMS Approval",        "ceo_contact"),
            "ISMS Manager":         ("isms_mgr_name",  "isms_mgr_dept", "Policy Implementation Audit",  "isms_mgr_contact"),
            "HR Head":              ("hr_name",        "hr_dept",       "Training & Awareness",         "hr_contact"),
            "System Administrator": ("sysadmin_name",  "sysadmin_dept", "IT security Management",       "sysadmin_contact"),
            "Technical Head":       ("tech_name",      "tech_dept",     "Security Implementation",      "tech_contact"),
        }
        for row in t.rows:
            role_cell = row.cells[0].text.strip() if row.cells else ""
            if role_cell in role_map:
                nk, dk, resp, ck = role_map[role_cell]
                if len(row.cells) > 1: _set_cell(row.cells[1], _v(d, nk))
                if len(row.cells) > 2: _set_cell(row.cells[2], _v(d, dk))
                if len(row.cells) > 3: _set_cell(row.cells[3], resp)
                if len(row.cells) > 4: _set_cell(row.cells[4], _v(d, ck))

    # ── 6. Table 5 — Scope Detail ─────────────────────────────────────────
    if len(doc.tables) > 5:
        t = doc.tables[5]
        label_to_value = {
            "Physical Location":        _v(d, "scope_location"),
            "Departments covered":      _v(d, "scope_depts"),
            "System Included":          _v(d, "scope_systems"),
            "Data Types (Client/Internal)": _v(d, "scope_data"),
            "Third Party Dependencies": _v(d, "scope_third_party"),
            "Scope Exclusions":         _v(d, "scope_exclusions"),
        }
        for row in t.rows:
            for ci, cell in enumerate(row.cells):
                stripped = cell.text.strip()
                for label, val in label_to_value.items():
                    if label in stripped and ci + 1 < len(row.cells):
                        _set_cell(row.cells[ci + 1], val)

    # ── 7. Table 6 — Roles & Responsibilities ────────────────────────────
    if len(doc.tables) > 6:
        t = doc.tables[6]
        cg_map = {
            "ISMS Manager":  ("isms_mgr_name", "Risk Management, Audit",    "High",   "System Admin"),
            "System Admin":  ("sysadmin_name",  "Access Control, Monitoring","Medium", "HR"),
            "HR":            ("hr_name",        "Training & Compliance",     "Medium", "Employees"),
            "Employees":     ("",               "Policy Adherence",          "Low",    ""),
        }
        for row in t.rows:
            role = row.cells[0].text.strip() if row.cells else ""
            if role in cg_map:
                nk, resp, auth, backup = cg_map[role]
                if len(row.cells) > 1: _set_cell(row.cells[1], _v(d, nk) if nk else "All Staff")
                if len(row.cells) > 2: _set_cell(row.cells[2], resp)
                if len(row.cells) > 3: _set_cell(row.cells[3], auth)
                if len(row.cells) > 4: _set_cell(row.cells[4], backup)

    # ── 8. Table 7 — ISMS Core Group ─────────────────────────────────────
    if len(doc.tables) > 7:
        t = doc.tables[7]
        # Cols: Sr. No. | Role | Name | Responsibility | Experience
        cg_rows = [
            ("1", "HR",           _v(d, "cg_hr_name"),  _v(d, "cg_hr_resp"),  _v(d, "cg_hr_exp")),
            ("2", "System Admin", _v(d, "cg_sa_name"),  _v(d, "cg_sa_resp"),  _v(d, "cg_sa_exp")),
            ("3", "Technical Head",_v(d,"cg_th_name"),  _v(d, "cg_th_resp"),  _v(d, "cg_th_exp")),
            ("4", "CEO",          _v(d, "cg_ceo_name"), _v(d, "cg_ceo_resp"), _v(d, "cg_ceo_exp")),
            ("5", "Auditor",      _v(d, "cg_aud_name"), _v(d, "cg_aud_resp"), _v(d, "cg_aud_exp")),
        ]
        data_rows = [row for row in t.rows if row.cells[0].text.strip() in {"1","2","3","4","5"}]
        for i, row in enumerate(data_rows):
            if i < len(cg_rows):
                sn, role, name, resp, exp = cg_rows[i]
                if len(row.cells) > 0: _set_cell(row.cells[0], sn)
                if len(row.cells) > 1: _set_cell(row.cells[1], role)
                if len(row.cells) > 2: _set_cell(row.cells[2], name)
                if len(row.cells) > 3: _set_cell(row.cells[3], resp)
                if len(row.cells) > 4: _set_cell(row.cells[4], exp)

    # ── 9. Table 8 — Contact Details ─────────────────────────────────────
    if len(doc.tables) > 8:
        t = doc.tables[8]
        contact_map = {
            "ISMS Manager":         ("contact_isms_name", "contact_isms_email", "contact_isms_phone"),
            "System Administrator": ("contact_sa_name",   "contact_sa_email",   "contact_sa_phone"),
            "HR":                   ("contact_hr_name",   "contact_hr_email",   "contact_hr_phone"),
        }
        for row in t.rows:
            role = row.cells[0].text.strip() if row.cells else ""
            if role in contact_map:
                nk, ek, pk = contact_map[role]
                if len(row.cells) > 1: _set_cell(row.cells[1], _v(d, nk))
                if len(row.cells) > 2: _set_cell(row.cells[2], _v(d, ek))
                if len(row.cells) > 3: _set_cell(row.cells[3], _v(d, pk))

    # ── 10. Table 9 — Escalation Matrix ──────────────────────────────────
    if len(doc.tables) > 9:
        t = doc.tables[9]
        esc_map = {
            "Level - 1": _v(d, "esc_l1_time"),
            "Level - 2": _v(d, "esc_l2_time"),
            "Level - 3": _v(d, "esc_l3_time"),
        }
        for row in t.rows:
            level = row.cells[0].text.strip() if row.cells else ""
            if level in esc_map:
                # Expected Response Time is last column
                if len(row.cells) > 3:
                    _set_cell(row.cells[3], esc_map[level])

    # ── 11. Table 11 — Interested Parties (static — already filled) ───────
    # Table 12 — Risk Register ─────────────────────────────────────────────
    if len(doc.tables) > 12:
        t = doc.tables[12]
        risks = d.get("risks", [])
        # Remove existing data rows (keep header)
        header_row = t.rows[0]
        while len(t.rows) > 1:
            tr = t.rows[-1]._tr
            tr.getparent().remove(tr)
        for idx, risk in enumerate(risks, start=1):
            impact = int(risk.get("risk_impact", 1) or 1)
            likelihood = int(risk.get("risk_likelihood", 1) or 1)
            risk_score = impact * likelihood
            risk_level = (
                "Critical" if risk_score >= 17 else
                "High"     if risk_score >= 13 else
                "Medium"   if risk_score >= 6  else "Low"
            )
            row = t.add_row()
            vals = [
                f"RISK-{idx:03d}",
                risk.get("risk_desc", "—"),
                str(risk.get("risk_impact", "—")),
                str(risk.get("risk_likelihood", "—")),
                str(risk_level),
                risk.get("risk_treatment", "—"),
                risk.get("risk_owner", "—"),
            ]
            for ci, val in enumerate(vals):
                if ci < len(row.cells):
                    _set_cell(row.cells[ci], val)

    # ── 12. Table 13 — Asset Register ────────────────────────────────────
    if len(doc.tables) > 13:
        t = doc.tables[13]
        assets = d.get("assets", [])
        while len(t.rows) > 1:
            tr = t.rows[-1]._tr
            tr.getparent().remove(tr)
        for idx, asset in enumerate(assets, start=1):
            row = t.add_row()
            vals = [
                asset.get("asset_name", "—"),
                asset.get("asset_type", "—"),
                asset.get("asset_owner", "—"),
                asset.get("asset_location", "—"),
                asset.get("asset_class", "—"),
            ]
            for ci, val in enumerate(vals):
                if ci < len(row.cells):
                    _set_cell(row.cells[ci], val)

    # ── 13. Table 14 — Policy Register ───────────────────────────────────
    if len(doc.tables) > 14:
        t = doc.tables[14]
        pol_map = {
            "ISMS Policy":   ("pol_isms_owner", "pol_isms_updated", "pol_isms_status"),
            "Access Policy": ("pol_acc_owner",  "pol_acc_updated",  "pol_acc_status"),
            "Backup Policy": ("pol_bak_owner",  "pol_bak_updated",  "pol_bak_status"),
        }
        for row in t.rows:
            pol = row.cells[0].text.strip() if row.cells else ""
            if pol in pol_map:
                ok, uk, sk = pol_map[pol]
                if len(row.cells) > 1: _set_cell(row.cells[1], _v(d, ok))
                if len(row.cells) > 3: _set_cell(row.cells[3], _v(d, uk))
                if len(row.cells) > 4: _set_cell(row.cells[4], _v(d, sk))

    # ── 14. Table 15 — Backup Register ───────────────────────────────────
    if len(doc.tables) > 15:
        t = doc.tables[15]
        backups = d.get("backups", [])
        while len(t.rows) > 1:
            tr = t.rows[-1]._tr
            tr.getparent().remove(tr)
        for backup in backups:
            row = t.add_row()
            vals = [
                backup.get("bak_system", "—"),
                backup.get("bak_freq", "—"),
                backup.get("bak_storage", "—"),
                backup.get("bak_responsible", "—"),
                backup.get("bak_tested", "—"),
            ]
            for ci, val in enumerate(vals):
                if ci < len(row.cells):
                    _set_cell(row.cells[ci], val)

    # ── 15. Table 16 — Technical Controls ────────────────────────────────
    if len(doc.tables) > 16:
        t = doc.tables[16]
        ctrl_map = {
            "Firewall":         ("ctrl_fw_impl",  "ctrl_fw_resp",  "ctrl_fw_remark"),
            "Antivirus":        ("ctrl_av_impl",  "ctrl_av_resp",  "ctrl_av_remark"),
            "Encryption":       ("ctrl_enc_impl", "ctrl_enc_resp", "ctrl_enc_remark"),
            "Patch Management": ("ctrl_pm_impl",  "ctrl_pm_resp",  "ctrl_pm_remark"),
        }
        for row in t.rows:
            ctrl = row.cells[0].text.strip() if row.cells else ""
            if ctrl in ctrl_map:
                ik, rk, rmk = ctrl_map[ctrl]
                if len(row.cells) > 1: _set_cell(row.cells[1], _v(d, ik))
                if len(row.cells) > 2: _set_cell(row.cells[2], _v(d, rk))
                if len(row.cells) > 3: _set_cell(row.cells[3], _v(d, rmk))

    # ── 16. Table 17 — Communication Matrix (static) ─────────────────────
    # ── 17. Table 18 — Approval Table ────────────────────────────────────
    if len(doc.tables) > 18:
        t = doc.tables[18]
        appr_map = {
            "Auditor":        _v(d, "appr_auditor"),
            "ISMS Manager":   _v(d, "appr_isms_mgr"),
            "CEO/ Director":  _v(d, "appr_ceo"),
        }
        for row in t.rows:
            role = row.cells[0].text.strip() if row.cells else ""
            if role in appr_map and len(row.cells) > 1:
                _set_cell(row.cells[1], appr_map[role])

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
