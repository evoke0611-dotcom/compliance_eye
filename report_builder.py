import io
import re
import datetime
import os
import tempfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFont
import matplotlib.pyplot as plt

def generate_logo(path):
    # Create a nice looking logo
    img = Image.new('RGB', (400, 100), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    # Draw shield icon
    d.polygon([(50, 15), (80, 15), (80, 70), (50, 55)], fill=(99, 102, 241))
    d.polygon([(20, 15), (50, 15), (50, 55), (20, 70)], fill=(139, 92, 246))
    
    # Text
    try:
        font = ImageFont.truetype("arial.ttf", 36)
        subfont = ImageFont.truetype("arial.ttf", 16)
    except IOError:
        font = ImageFont.load_default()
        subfont = ImageFont.load_default()
        
    d.text((95, 20), "ComplianceEye", fill=(30, 27, 75), font=font)
    d.text((100, 65), "OFFICIAL ISO 27001 AUDIT", fill=(99, 102, 241), font=subfont)
    img.save(path)

def add_charts_to_doc(doc, stats):
    if not stats:
        return
        
    doc.add_page_break()
    doc.add_heading("📊 Assessment Insights & Analytics", level=1)
    
    # Gap Analysis Pie Chart
    gap_data = stats.get("gap_counts", {"Critical": 0, "Partial": 0, "Compliant": 0, "Best Practice": 0})
    labels = list(gap_data.keys())
    values = list(gap_data.values())
    colors = ['#ef4444', '#f59e0b', '#3b82f6', '#10b981']
    
    plt.figure(figsize=(6, 4))
    plt.pie(values, labels=labels, autopct='%1.1f%%', startangle=140, colors=colors)
    plt.title('Gap Analysis')
    plt.axis('equal')
    
    gap_img = io.BytesIO()
    plt.savefig(gap_img, format='png', bbox_inches='tight')
    plt.close()
    gap_img.seek(0)
    
    doc.add_paragraph("The following chart illustrates the distribution of identified gaps across the assessment bank.")
    doc.add_picture(gap_img, width=Inches(5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    # Risk Severity Bar Chart
    risk_data = stats.get("risk_counts", {"Critical": 0, "High": 0, "Medium": 0, "Low": 0})
    r_labels = list(risk_data.keys())
    r_values = list(risk_data.values())
    r_colors = ['#b91c1c', '#dc2626', '#f59e0b', '#3b82f6']
    
    plt.figure(figsize=(6, 4))
    plt.bar(r_labels, r_values, color=r_colors)
    plt.title('Risk Severity Distribution')
    plt.ylabel('Number of Controls')
    
    risk_img = io.BytesIO()
    plt.savefig(risk_img, format='png', bbox_inches='tight')
    plt.close()
    risk_img.seek(0)
    
    doc.add_paragraph("The risk distribution highlights the severity of identified security risks.")
    doc.add_picture(risk_img, width=Inches(5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

def create_docx_report(markdown_text, stats=None):
    doc = Document()
    
    # --- STYLES ---
    styles = doc.styles
    style = styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)
    
    if 'Title' in styles:
        title_style = styles['Title']
        title_style.font.name = 'Arial'
        title_style.font.size = Pt(26)
        title_style.font.color.rgb = RGBColor(30, 27, 75)
        title_style.font.bold = True
        
    # --- HEADER WITH LOGO ---
    section = doc.sections[0]
    header = section.header
    header_para = header.paragraphs[0]
    logo_path = "temp_logo.png"
    try:
        generate_logo(logo_path)
        run = header_para.add_run()
        run.add_picture(logo_path, width=Inches(2.5))
        header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    except Exception as e:
        print("Logo generation failed", e)
        header_para.text = "ComplianceEye - ISO 27001 AI Auditor"
        header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
    # --- COVER PAGE ---
    for _ in range(5):
        doc.add_paragraph()
    title = doc.add_heading("ISO 27001 Information Security Documentation", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    subtitle = doc.add_paragraph("Official Audit Report")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(16)
    subtitle.runs[0].font.color.rgb = RGBColor(100, 116, 139)
    doc.add_paragraph()
    date_p = doc.add_paragraph(f"Generated on: {datetime.datetime.now().strftime('%Y-%m-%d')}")
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    
    # --- PARSE MARKDOWN ---
    # Filter out initial chatter before the first heading
    start_idx = markdown_text.find('#')
    if start_idx != -1:
        markdown_text = markdown_text[start_idx:]
        
    lines = markdown_text.split('\n')
    
    for line in lines:
        line = line.strip()
        if not line or line == '---':
            continue
            
        clean_line = line.replace('**', '')
        
        # Skip the main title since we put it on the cover page
        if clean_line.startswith('# ') and "ISO 27001" in clean_line:
            continue
            
        if clean_line.startswith('# '):
            h = doc.add_heading(clean_line.replace('# ', '').strip(), level=1)
            h.runs[0].font.color.rgb = RGBColor(79, 70, 229)
        elif clean_line.startswith('## '):
            h = doc.add_heading(clean_line.replace('## ', '').strip(), level=2)
            h.runs[0].font.color.rgb = RGBColor(99, 102, 241)
        elif clean_line.startswith('### '):
            h = doc.add_heading(clean_line.replace('### ', '').strip(), level=3)
            h.runs[0].font.color.rgb = RGBColor(30, 27, 75)
        elif clean_line.startswith('- '):
            item_text = clean_line.replace('- ', '', 1).strip()
            # If it's a key-value pair, make the key bold
            if ':' in item_text:
                parts = item_text.split(':', 1)
                p = doc.add_paragraph(style='List Bullet')
                r_key = p.add_run(parts[0].strip() + ':')
                r_key.bold = True
                p.add_run(' ' + parts[1].strip())
            else:
                doc.add_paragraph(item_text, style='List Bullet')
        else:
            if "Sure!" in clean_line or "Here's the" in clean_line or "Here is the" in clean_line:
                continue
            doc.add_paragraph(clean_line)
            
    # --- ADD CHARTS ---
    if stats:
        add_charts_to_doc(doc, stats)
            
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _extract_option_letter(selected_text):
    if not selected_text:
        return ""
    match = re.match(r"^\s*([A-Da-d])\.", str(selected_text).strip())
    if match:
        return match.group(1).upper()
    return ""


def _safe_pdf_text(value, max_word_len=40):
    text = str(value or "-")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[^\x09\x0A\x0D\x20-\x7E]", " ", text)
    text = re.sub(r"\s+", " ", text).strip() or "-"
    parts = []
    for token in text.split(" "):
        if len(token) <= max_word_len:
            parts.append(token)
        else:
            parts.extend(token[i:i + max_word_len] for i in range(0, len(token), max_word_len))
    return " ".join(parts)


def _safe_multicell(pdf, h, text):
    content = _safe_pdf_text(text)
    available_w = pdf.w - pdf.l_margin - pdf.r_margin
    if available_w <= 1:
        available_w = 180
    pdf.multi_cell(available_w, h, content)


def _build_assessment_summary(assessment_answers, clause_metadata):
    meta_map = {}
    for item in clause_metadata or []:
        cid = str(item.get("control_id", "")).strip()
        if not cid:
            continue
        if cid not in meta_map:
            meta_map[cid] = item
        elif not meta_map[cid].get("description") and item.get("description"):
            meta_map[cid] = item

    processed = []
    gap_counts = {"Critical Gap": 0, "Partial Gap": 0, "No Gap": 0, "Best Practice": 0}
    risk_counts = {"Critical": 0, "High": 0, "Medium": 0, "Low": 0}

    for ans in assessment_answers or []:
        letter = _extract_option_letter(ans.get("selected", ""))
        status_map = {"A": "Critical Gap", "B": "Partial Gap", "C": "No Gap", "D": "Best Practice"}
        maturity_map = {"A": 1, "B": 2, "C": 3, "D": 4}
        risk_map = {"A": "Critical", "B": "High", "C": "Medium", "D": "Low"}
        status = status_map.get(letter, "Partial Gap")
        maturity = maturity_map.get(letter, 2)
        risk = risk_map.get(letter, "High")
        gap_counts[status] = gap_counts.get(status, 0) + 1
        risk_counts[risk] = risk_counts.get(risk, 0) + 1
        clause = str(ans.get("clause", "")).strip()
        meta = meta_map.get(clause, {})
        processed.append(
            {
                "clause": clause or "N/A",
                "question": ans.get("question", ""),
                "selected": ans.get("selected", ""),
                "gap_status": status,
                "maturity_score": maturity,
                "risk_level": risk,
                "title": meta.get("title", "ISO 27001 Clause"),
                "description": meta.get("description", ""),
            }
        )
    return processed, gap_counts, risk_counts


def create_structured_docx_report(
    basic_info,
    isms_answers,
    gap_assessment_answers,
    risk_assessment_answers,
    clause_metadata,
    generated_on=None,
):
    generated_on = generated_on or datetime.datetime.now().strftime("%Y-%m-%d")
    doc = Document()

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)

    def _extract_letter(selected_text: str):
        s = str(selected_text or "").strip().upper()
        m = re.match(r"^\s*([A-E])[.)]\s*", s)
        if m:
            return m.group(1)
        if s[:1] in {"A", "B", "C", "D", "E"}:
            return s[:1]
        return ""

    STOPWORDS = {
        "the", "and", "or", "for", "with", "into", "to", "of", "in", "on", "at", "by",
        "is", "are", "be", "as", "an", "it", "that", "this", "from", "are", "their",
        "across", "how", "what", "which", "when", "where", "about", "your", "not",
        "can", "could", "should", "would", "may", "will", "must", "do", "does",
        "does", "likely", "possible", "unlikely", "almost", "very", "minimal",
        "minor", "moderate", "significant", "severe", "risk",
    }

    def _tokenize(s: str):
        tokens = re.findall(r"[a-z0-9]+", (s or "").lower())
        return {t for t in tokens if len(t) >= 4 and t not in STOPWORDS}

    def _gap_from_letter(letter: str):
        # Per Gap & Risk Logic.pdf: maturity score is 1–5.
        gap_status_map = {
            "A": ("Critical Gap", 1),
            "B": ("Major Gap", 2),
            "C": ("Partial Gap", 3),
            "D": ("Minor Gap", 4),
            "E": ("No Gap", 5),
        }
        return gap_status_map.get(letter, ("Partial Gap", 3))

    def _risk_from_letter(letter: str):
        # Per Risk Assessment MCQs.pdf: A–E represent (likelihood, impact) = 1..5.
        n = {"A": 1, "B": 2, "C": 3, "D": 4, "E": 5}.get(letter, 3)
        l = n
        i = n
        score = l * i
        # Risk levels inferred from the provided example scores.
        if score >= 17:
            level = "Critical"
        elif score >= 13:
            level = "High"
        elif score >= 6:
            level = "Medium"
        else:
            level = "Low"
        return l, i, score, level

    def _risk_suggestions(level: str, clause_title: str):
        if level == "Critical":
            return (
                f"Immediate: Contain/mitigate {clause_title} gaps and protect high-value assets.",
                f"Short-term: Implement reinforced controls + monitoring evidence for {clause_title}.",
                f"Long-term: Integrate into enterprise risk treatment and continuously improve."
            )
        if level == "High":
            return (
                f"Immediate: Strengthen controls aligned to {clause_title}.",
                f"Short-term: Implement risk treatment plan and validate effectiveness.",
                f"Long-term: Run regular reviews and refine according to results."
            )
        if level == "Medium":
            return (
                f"Immediate: Document risks and assign ownership for {clause_title}.",
                f"Short-term: Plan control improvements and start evidence collection.",
                f"Long-term: Monitor trends and update risk treatment actions."
            )
        return (
            f"Immediate: Maintain current controls for {clause_title}.",
            f"Short-term: Ensure periodic review and evidence readiness.",
            f"Long-term: Drive continual improvement and maturity."
        )

    def _gap_suggestions(status: str, clause_title: str):
        if status == "Critical Gap":
            return (
                f"Immediate: Define and approve required governance for {clause_title}.",
                f"Short-term: Implement controls/processes and establish evidence.",
                f"Long-term: Measure effectiveness and embed continuous improvement."
            )
        if status == "Major Gap":
            return (
                f"Immediate: Document missing requirements for {clause_title}.",
                f"Short-term: Implement controls across the in-scope environment.",
                f"Long-term: Validate via audits and track improvement actions."
            )
        if status == "Partial Gap":
            return (
                f"Immediate: Close documentation/evidence gaps for {clause_title}.",
                f"Short-term: Standardize the process and verify implementation.",
                f"Long-term: Improve maturity through monitoring and reviews."
            )
        if status == "Minor Gap":
            return (
                f"Immediate: Address smaller gaps for {clause_title} and correct exceptions.",
                f"Short-term: Improve consistency of controls and evidence.",
                f"Long-term: Keep maturing through periodic reviews."
            )
        return (
            f"Immediate: Maintain compliance for {clause_title}.",
            f"Short-term: Ensure evidence remains current and complete.",
            f"Long-term: Pursue continual improvement for maturity."
        )

    # Clause matching from iso27001_metadata.json
    clause_index = []
    for item in clause_metadata or []:
        control_id = str(item.get("control_id", "")).strip()
        title = str(item.get("title", "")).strip()
        desc = str(item.get("description", "")).strip()
        text = f"{control_id} {title} {desc[:900]}"
        clause_index.append(
            {
                "control_id": control_id,
                "title": title,
                "description": desc,
                "tokens": _tokenize(text),
            }
        )

    def _best_clause(question_text: str):
        q_tokens = _tokenize(question_text)
        best = None
        best_score = -1
        for c in clause_index:
            score = len(q_tokens & c["tokens"])
            if score > best_score:
                best_score = score
                best = c
        if not best:
            return {"control_id": "N/A", "title": "N/A", "description": ""}
        return best

    logo_path = os.path.join(tempfile.gettempdir(), "complianceeye_logo_docx.png")
    try:
        generate_logo(logo_path)
    except Exception:
        logo_path = None

    # Header logo on every page
    if logo_path and os.path.exists(logo_path):
        for section in doc.sections:
            section.different_first_page_header_footer = False
            header = section.header
            para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = para.add_run()
            run.add_picture(logo_path, width=Inches(2.1))

    # Cover page
    doc.add_paragraph()
    if logo_path and os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.add_run().add_picture(logo_path, width=Inches(3.4))
    title = doc.add_heading("ISO 27001 Formal Audit Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("ComplianceEye Audit Documentation")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    doc.add_paragraph(f"Reporting Date: {generated_on}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Company Name: {basic_info.get('org_name', '-')}" ).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Auditor Name: {basic_info.get('emp_name', '-')}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # Step 1: Basic info
    doc.add_heading("1. Company and Auditor Information", level=1)
    fields = [
        ("Organization Name", basic_info.get("org_name", "-")),
        ("Business Nature", basic_info.get("business_nature", "-")),
        ("Services / Products", basic_info.get("services", "-")),
        ("Information Security Mission", basic_info.get("is_mission", "-")),
        ("Information Security Importance", basic_info.get("is_importance", "-")),
        ("Key Stakeholders", basic_info.get("stakeholders", "-")),
        ("Stakeholder Expectations", basic_info.get("stakeholder_exp", "-")),
        ("Auditor Name", basic_info.get("emp_name", "-")),
        ("Auditor Role", basic_info.get("emp_role", "-")),
        ("Auditor Email", basic_info.get("emp_email", "-")),
        ("Auditor Contact", basic_info.get("emp_contact", "-")),
        ("ISMS Responsible", basic_info.get("isms_responsible", "-")),
        ("ISMS Core Group", basic_info.get("isms_core_group", "-")),
        ("Management Commitment", basic_info.get("mgmt_commitment", "-")),
    ]
    for label, value in fields:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{label}: ").bold = True
        p.add_run(str(value or "-"))

    # Step 2: ISMS policy Q&A
    doc.add_page_break()
    doc.add_heading("2. ISMS Policy Questions and Answers", level=1)
    if not isms_answers:
        doc.add_paragraph("No ISMS policy responses available.")
    else:
        for idx, qa in enumerate(isms_answers, start=1):
            q = doc.add_paragraph()
            q.add_run(f"Q{idx}. {qa.get('question', '')}").bold = True
            doc.add_paragraph(f"Answer: {qa.get('answer', '-')}")

    # Step 3: Gap Assessment
    gap_answers = gap_assessment_answers or []
    gap_results = []
    gap_counts = {"Critical Gap": 0, "Major Gap": 0, "Partial Gap": 0, "Minor Gap": 0, "No Gap": 0}
    gap_score_total = 0
    gap_score_counts = {1: 0, 2: 0, 3: 0, 4: 0, 5: 0}

    for idx, ans in enumerate(gap_answers, start=1):
        selected = ans.get("selected", "")
        letter = _extract_letter(selected)
        status, score = _gap_from_letter(letter)
        gap_score_total += score
        gap_score_counts[score] = gap_score_counts.get(score, 0) + 1
        gap_counts[status] = gap_counts.get(status, 0) + 1

        question_text = ans.get("question", "") or ""
        matched_clause = _best_clause(question_text)
        clause_title = matched_clause.get("title") or matched_clause.get("control_id") or "N/A"
        immediate, short_term, long_term = _gap_suggestions(status, clause_title)
        action_plan = f"{immediate} {short_term} {long_term}"

        gap_results.append(
            {
                "q_no": idx,
                "question": question_text,
                "selected": selected,
                "letter": letter,
                "gap_status": status,
                "score": score,
                "clause_id": matched_clause.get("control_id", "N/A"),
                "clause_title": clause_title,
                "action_plan": action_plan,
            }
        )

    gap_maturity_percentage = 0.0
    if gap_results:
        gap_maturity_percentage = (gap_score_total / (5 * len(gap_results))) * 100.0

    doc.add_page_break()
    doc.add_heading("3. Gap Assessment", level=1)
    doc.add_paragraph(f"Overall Gap Maturity: {gap_maturity_percentage:.0f}%")

    total_gaps_identified = sum(gap_counts[s] for s in ["Critical Gap", "Major Gap", "Partial Gap", "Minor Gap"])
    doc.add_paragraph(f"Total Gaps Identified: {total_gaps_identified}")

    # Charts (per Gap & Risk Logic.pdf)
    gap_img = io.BytesIO()
    plt.figure(figsize=(6, 4))
    gap_labels = ["Critical Gap", "Major Gap", "Partial Gap", "Minor Gap", "No Gap"]
    gap_vals = [gap_counts.get(k, 0) for k in gap_labels]
    plt.pie(
        gap_vals,
        labels=gap_labels,
        autopct="%1.1f%%",
        startangle=140,
        colors=["#ef4444", "#f59e0b", "#10b981", "#3b82f6", "#22c55e"],
    )
    plt.title("Gap Distribution")
    plt.tight_layout()
    plt.savefig(gap_img, format="png", dpi=120)
    plt.close()
    gap_img.seek(0)
    doc.add_picture(gap_img, width=Inches(5.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    bar_img = io.BytesIO()
    plt.figure(figsize=(6, 4))
    plt.bar([1, 2, 3, 4, 5], [gap_score_counts[i] for i in [1, 2, 3, 4, 5]], color=["#ef4444", "#f59e0b", "#f59e0b", "#10b981", "#22c55e"])
    plt.xticks([1, 2, 3, 4, 5])
    plt.title("Gap Maturity Scores")
    plt.xlabel("Maturity Score (1–5)")
    plt.ylabel("Number of Requirements")
    plt.tight_layout()
    plt.savefig(bar_img, format="png", dpi=120)
    plt.close()
    bar_img.seek(0)
    doc.add_picture(bar_img, width=Inches(5.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Gap findings table
    doc.add_paragraph("Gap Assessment Findings (MCQ → Maturity Score → Suggested Actions)")
    gap_table = doc.add_table(rows=1, cols=6)
    gap_table.style = "Table Grid"
    hdr = gap_table.rows[0].cells
    hdr[0].text = "Q#"
    hdr[1].text = "Matched Clause"
    hdr[2].text = "Selected"
    hdr[3].text = "Gap Status"
    hdr[4].text = "Score"
    hdr[5].text = "Action Plan (Immediate/Short/Long)"

    for r in gap_results:
        row_cells = gap_table.add_row().cells
        row_cells[0].text = str(r["q_no"])
        row_cells[1].text = f"{r['clause_id']} - {r['clause_title']}"
        row_cells[2].text = str(r["selected"])
        row_cells[3].text = r["gap_status"]
        row_cells[4].text = str(r["score"])
        row_cells[5].text = r["action_plan"]

    # Step 4: Risk Assessment
    risk_answers = risk_assessment_answers or []
    risk_results = []
    risk_counts = {"Low": 0, "Medium": 0, "High": 0, "Critical": 0}
    risk_score_total = 0
    risk_matrix_counts = [[0 for _ in range(5)] for __ in range(5)]  # likelihood(1-5) x impact(1-5)
    risk_matrix_level = [[0 for _ in range(5)] for __ in range(5)]  # numeric level

    level_to_num = {"Low": 1, "Medium": 2, "High": 3, "Critical": 4}

    for idx, ans in enumerate(risk_answers, start=1):
        selected = ans.get("selected", "")
        letter = _extract_letter(selected)
        l, i, score, level = _risk_from_letter(letter)
        risk_score_total += score
        risk_counts[level] = risk_counts.get(level, 0) + 1

        question_text = ans.get("question", "") or ""
        matched_clause = _best_clause(question_text)
        clause_title = matched_clause.get("title") or matched_clause.get("control_id") or "N/A"
        immediate, short_term, long_term = _risk_suggestions(level, clause_title)
        action_plan = f"{immediate} {short_term} {long_term}"

        risk_results.append(
            {
                "q_no": idx,
                "question": question_text,
                "selected": selected,
                "letter": letter,
                "clause_id": matched_clause.get("control_id", "N/A"),
                "clause_title": clause_title,
                "likelihood": l,
                "impact": i,
                "risk_score": score,
                "risk_level": level,
                "action_plan": action_plan,
            }
        )

        li = l - 1
        ii = i - 1
        risk_matrix_counts[li][ii] += 1
        risk_matrix_level[li][ii] = max(risk_matrix_level[li][ii], level_to_num.get(level, 1))

    overall_risk_exposure = 0.0
    if risk_results:
        overall_risk_exposure = (risk_score_total / (25 * len(risk_results))) * 100.0

    doc.add_page_break()
    doc.add_heading("4. Risk Assessment", level=1)
    doc.add_paragraph(f"Overall Risk Exposure: {overall_risk_exposure:.0f}%")

    high_risks_count = sum(risk_counts.get(k, 0) for k in ["High", "Critical"])
    doc.add_paragraph(f"High Risks: {high_risks_count}")

    # Heatmap (likelihood × impact)
    heat_img = io.BytesIO()
    import numpy as _np  # local import to avoid changing module-level deps
    level_mat = _np.array(risk_matrix_level, dtype=float)
    count_mat = _np.array(risk_matrix_counts, dtype=float)

    plt.figure(figsize=(7, 5))
    plt.imshow(level_mat, cmap="YlOrRd", interpolation="nearest", vmin=0, vmax=4)
    plt.colorbar(label="Risk Level (1–4)")
    plt.xticks(range(5), [1, 2, 3, 4, 5])
    plt.yticks(range(5), [1, 2, 3, 4, 5])
    plt.xlabel("Impact (1–5)")
    plt.ylabel("Likelihood (1–5)")
    plt.title("Risk Heatmap (Likelihood × Impact)")

    for yi in range(5):
        for xi in range(5):
            cnt = int(risk_matrix_counts[yi][xi])
            if cnt <= 0:
                continue
            lv = int(risk_matrix_level[yi][xi])
            plt.text(xi, yi, f"{lv}\n({cnt})", ha="center", va="center", fontsize=8, color="black")

    plt.tight_layout()
    plt.savefig(heat_img, format="png", dpi=120)
    plt.close()
    heat_img.seek(0)
    doc.add_picture(heat_img, width=Inches(5.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Risk register table
    doc.add_paragraph("Risk Register (Likelihood × Impact → Score → Level)")
    risk_table = doc.add_table(rows=1, cols=7)
    risk_table.style = "Table Grid"
    hdr = risk_table.rows[0].cells
    hdr[0].text = "Matched Clause"
    hdr[1].text = "Threat (MCQ Question)"
    hdr[2].text = "Likelihood"
    hdr[3].text = "Impact"
    hdr[4].text = "Score"
    hdr[5].text = "Risk Level"
    hdr[6].text = "Action Plan (Immediate/Short/Long)"

    for rr in risk_results:
        row_cells = risk_table.add_row().cells
        row_cells[0].text = f"{rr['clause_id']} - {rr['clause_title']}"
        row_cells[1].text = str(rr["question"])[:250]
        row_cells[2].text = str(rr["likelihood"])
        row_cells[3].text = str(rr["impact"])
        row_cells[4].text = str(rr["risk_score"])
        row_cells[5].text = rr["risk_level"]
        row_cells[6].text = rr["action_plan"]

    # Step 5: Executive Summary
    doc.add_page_break()
    doc.add_heading("5. Executive Summary", level=1)
    doc.add_paragraph(f"Gap Maturity: {gap_maturity_percentage:.0f}%")
    doc.add_paragraph(f"Risk Exposure: {overall_risk_exposure:.0f}%")
    doc.add_paragraph(
        "This audit report translates questionnaire answers into measurable maturity and risk indicators, "
        "and provides improvement actions aligned with ISO/IEC 27001 clause requirements."
    )

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out


def create_structured_pdf_report(
    basic_info,
    isms_answers,
    assessment_answers,
    clause_metadata,
    generated_on=None,
):
    generated_on = generated_on or datetime.datetime.now().strftime("%Y-%m-%d")
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.set_margins(12, 12, 12)

    processed, gap_counts, risk_counts = _build_assessment_summary(assessment_answers, clause_metadata)

    # Cover Page: basic information
    pdf.add_page()
    pdf.set_font("Arial", "B", 18)
    pdf.cell(0, 12, "ISO 27001 Final Audit Report", ln=True, align="C")
    pdf.ln(2)
    pdf.set_font("Arial", "", 11)
    pdf.cell(0, 8, f"Generated On: {generated_on}", ln=True, align="C")
    pdf.ln(8)

    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 9, "Company and Auditor Information", ln=True)
    pdf.set_font("Arial", "", 11)
    fields = [
        ("Organization Name", basic_info.get("org_name", "-")),
        ("Business Nature", basic_info.get("business_nature", "-")),
        ("Services / Products", basic_info.get("services", "-")),
        ("Information Security Mission", basic_info.get("is_mission", "-")),
        ("Employee Name", basic_info.get("emp_name", "-")),
        ("Employee Role", basic_info.get("emp_role", "-")),
        ("Employee Email", basic_info.get("emp_email", "-")),
        ("Employee Contact", basic_info.get("emp_contact", "-")),
        ("ISMS Responsible", basic_info.get("isms_responsible", "-")),
    ]
    for label, value in fields:
        pdf.set_font("Arial", "B", 11)
        _safe_multicell(pdf, 7, f"{label}:")
        pdf.set_font("Arial", "", 11)
        _safe_multicell(pdf, 7, value)
        pdf.ln(1)

    # Page 2: ISMS policy questions and answers
    pdf.add_page()
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 9, "ISMS Policy - Questions and Responses", ln=True)
    pdf.ln(2)
    if not isms_answers:
        pdf.set_font("Arial", "", 11)
        _safe_multicell(pdf, 7, "No ISMS policy responses were captured for this session.")
    else:
        for idx, qa in enumerate(isms_answers, start=1):
            pdf.set_font("Arial", "B", 11)
            _safe_multicell(pdf, 7, f"Q{idx}. {qa.get('question', '')}")
            pdf.set_font("Arial", "", 11)
            _safe_multicell(pdf, 7, f"Answer: {qa.get('answer', '-')}")
            pdf.ln(2)

    # Page 3+: Gap analysis detailed section
    pdf.add_page()
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 9, "Risk and Gap Analysis", ln=True)
    pdf.ln(2)
    pdf.set_font("Arial", "", 10)
    _safe_multicell(
        pdf,
        6,
        "This section shows clause-wise gaps, selected answers, ISO reference details, and practical recommendations.",
    )
    pdf.ln(2)

    if not processed:
        pdf.set_font("Arial", "", 11)
        _safe_multicell(pdf, 7, "No gap assessment responses were captured for this session.")
    else:
        for idx, item in enumerate(processed, start=1):
            if pdf.get_y() > 245:
                pdf.add_page()
            recommendation = "Maintain and continuously improve this control."
            if item["gap_status"] in ("Critical Gap", "Partial Gap"):
                recommendation = "Define a documented control, assign ownership, implement it, and review effectiveness regularly."

            pdf.set_font("Arial", "B", 11)
            _safe_multicell(pdf, 7, f"{idx}. Clause {item['clause']} - {item['title']}")
            pdf.set_font("Arial", "", 10)
            _safe_multicell(pdf, 6, f"Question: {item['question']}")
            _safe_multicell(pdf, 6, f"Given Answer: {item['selected']}")
            _safe_multicell(pdf, 6, f"Gap Status: {item['gap_status']} | Maturity Score: {item['maturity_score']}/4 | Risk: {item['risk_level']}")
            if item["description"]:
                _safe_multicell(pdf, 6, f"ISO Clause Detail: {item['description'][:420]}")
            _safe_multicell(pdf, 6, f"What should be done: {recommendation}")
            pdf.ln(2)

    # Charts and summary page
    pdf.add_page()
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 9, "Summary and Charts", ln=True)
    pdf.ln(2)

    # Build chart images
    gap_fig = io.BytesIO()
    plt.figure(figsize=(5.5, 3.2))
    plt.pie(
        list(gap_counts.values()),
        labels=list(gap_counts.keys()),
        autopct="%1.1f%%",
        startangle=140,
        colors=["#ef4444", "#f59e0b", "#10b981", "#3b82f6"],
    )
    plt.title("Gap Distribution")
    plt.tight_layout()
    plt.savefig(gap_fig, format="png", dpi=120)
    plt.close()
    gap_fig.seek(0)

    risk_fig = io.BytesIO()
    plt.figure(figsize=(5.5, 3.2))
    plt.bar(
        list(risk_counts.keys()),
        list(risk_counts.values()),
        color=["#b91c1c", "#dc2626", "#f59e0b", "#3b82f6"],
    )
    plt.title("Risk Severity Distribution")
    plt.tight_layout()
    plt.savefig(risk_fig, format="png", dpi=120)
    plt.close()
    risk_fig.seek(0)

    temp_paths = []
    try:
        for fig in (gap_fig, risk_fig):
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
            tmp.write(fig.getvalue())
            tmp.flush()
            tmp.close()
            temp_paths.append(tmp.name)

        pdf.set_font("Arial", "", 10)
        _safe_multicell(
            pdf,
            6,
            f"Total controls assessed: {len(processed)} | Critical+Partial gaps: {gap_counts['Critical Gap'] + gap_counts['Partial Gap']}",
        )
        pdf.ln(2)
        pdf.image(temp_paths[0], x=20, y=pdf.get_y(), w=170)
        pdf.ln(78)
        pdf.image(temp_paths[1], x=20, y=pdf.get_y(), w=170)
        pdf.ln(78)
    finally:
        for p in temp_paths:
            try:
                os.remove(p)
            except OSError:
                pass

    pdf.set_font("Arial", "B", 11)
    _safe_multicell(pdf, 7, "Executive Summary")
    pdf.set_font("Arial", "", 10)
    critical_partial = gap_counts["Critical Gap"] + gap_counts["Partial Gap"]
    summary = (
        f"The organization has {critical_partial} controls that need immediate or near-term improvement. "
        f"Maturity improves as controls move from option A/B to C/D with documented process ownership, "
        "periodic review, and measurable ISMS governance."
    )
    _safe_multicell(pdf, 6, summary)

    output = pdf.output(dest="S")
    if isinstance(output, str):
        return output.encode("latin-1")
    return bytes(output)
